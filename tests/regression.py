from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def count_non_empty(rows: list[dict], field: str) -> int:
    return sum(1 for row in rows if row.get(field))


def load_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(x).strip() if x is not None else "" for x in rows[0]]

    records = []
    for row in rows[1:]:
        records.append(dict(zip(headers, row)))

    return records


def check_metric(name: str, actual: int, minimum: int) -> bool:
    ok = actual >= minimum
    status = "PASS" if ok else "FAIL"
    print(f"{name:<20} {status:<5} actual={actual:<5} expected>={minimum}")
    return ok


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage:")
        print("  python tests/regression.py <output.xlsx> <expected.json>")
        return 2

    xlsx_path = Path(sys.argv[1])
    expected_path = Path(sys.argv[2])

    rows = load_xlsx(xlsx_path)
    expected = json.loads(expected_path.read_text(encoding="utf-8"))

    metrics = {
        "records": len(rows),
        "emails": count_non_empty(rows, "email"),
        "phones": count_non_empty(rows, "phone"),
        "websites": count_non_empty(rows, "website"),
        "facebook": count_non_empty(rows, "facebook"),
        "linkedin": count_non_empty(rows, "linkedin"),
        "instagram": count_non_empty(rows, "instagram"),
        "youtube": count_non_empty(rows, "youtube"),
        "twitter": count_non_empty(rows, "twitter"),
    }

    mapping = {
        "records": "minimum_records",
        "emails": "minimum_emails",
        "phones": "minimum_phones",
        "websites": "minimum_websites",
        "facebook": "minimum_facebook",
        "linkedin": "minimum_linkedin",
        "instagram": "minimum_instagram",
        "youtube": "minimum_youtube",
        "twitter": "minimum_twitter",
    }

    print("\n====================================")
    print("UBDI REGRESSION TEST")
    print("====================================\n")

    results = []

    for metric_name, expected_key in mapping.items():
        results.append(
            check_metric(
                metric_name,
                metrics[metric_name],
                int(expected.get(expected_key, 0)),
            )
        )

    print("\n====================================")
    print("RESULT:", "PASS" if all(results) else "FAIL")
    print("====================================\n")

    return 0 if all(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())