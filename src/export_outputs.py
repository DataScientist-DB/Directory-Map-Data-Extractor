from __future__ import annotations

import csv
import json
import re
import html
from pathlib import Path
from typing import Any, Dict, List, Optional
from src.taxonomy_static import map_product_labels_to_rpc, RPC_CATEGORY_MAP

from openpyxl import Workbook

# =========================
# Taxonomy decoding helpers
# =========================

_SPLIT_CODES_RE = re.compile(r"[,\t;|]+")  # comma, tab, semicolon, pipe
_BAD_LABELS = {"…", "...", "⋯", "·", "-"}


def _clean_text(v: Any) -> str:
    if v is None:
        return ""
    s = html.unescape(str(v))
    s = " ".join(s.split())
    return s.strip()


def _normalize_code(code: str) -> str:
    """
    Normalize a taxonomy code:
      - trim
      - lowercase
      - convert ssN -> rssN (site sometimes emits ss1)
    """
    c = (code or "").strip().lower()
    if not c:
        return ""
    if re.fullmatch(r"ss\d+", c):
        return "rss" + c[2:]
    return c


def _parse_codes_any(v: Any) -> list[str]:
    """
    Accepts:
      - "rpc1, rpc2" or "rpc1; rpc2" or "rpc1\t rpc2" or "rpc1|rpc2"
      - ["rpc1", "rpc2"]
      - "" / None
    Returns unique list preserving order.
    """
    if v is None:
        return []
    if isinstance(v, list):
        raw = [str(x) for x in v]
    else:
        raw = _SPLIT_CODES_RE.split(str(v))

    out: list[str] = []
    seen: set[str] = set()
    for part in raw:
        code = _normalize_code(part)
        if not code:
            continue
        if code not in seen:
            seen.add(code)
            out.append(code)
    return out


def _normalize_tax_map(m: Dict[str, str]) -> Dict[str, str]:
    """
    Normalize mapping keys to lowercase, clean labels, and drop placeholders like "…".
    """
    out: Dict[str, str] = {}
    for k, v in (m or {}).items():
        kk = _normalize_code(str(k))
        if not kk:
            continue
        vv = _clean_text(v)
        if not vv or vv in _BAD_LABELS:
            continue
        out[kk] = vv
    return out


def _decode_codes(codes: list[str], mapping: dict[str, str]) -> list[str]:
    """
    Returns decoded names for codes.
    Falls back to original code if unknown.
    Ensures uniqueness preserving order.
    """
    decoded: list[str] = []
    seen: set[str] = set()

    for raw in (codes or []):
        c = _normalize_code(raw)
        if not c:
            continue

        # accept only rpc/rss forms
        if not re.fullmatch(r"(rpc|rss)\d+", c):
            continue

        name = mapping.get(c)
        val = _clean_text(name) if name else c

        if val and val not in seen:
            seen.add(val)
            decoded.append(val)

    return decoded


# =========================
# Output column definitions
# =========================

DEFAULT_COLUMNS = [
    "entity_name",
    "category_names",
    "service_names",
    "products",
    "products_rpc_codes",
    "products_rpc_names",

    "email",
    "phone",
    "fax",
    "website",

    "linkedin",
    "facebook",
    "instagram",
    "youtube",
    "twitter",

    "location",
    "address",
    "city",
    "state",
    "postal_code",

    "description",
    "hours",
    "driving_directions",

    "profile_url",
    "source_url",
    "architecture",
    "crawl_mode",
    "confidence_score",

    "status",
    "records_found",
    "blocked_reason",

    "website_enrichment_status",
    "website_enrichment_error",

    # Business Intelligence
    "intelligence_score",
    "intelligence_grade",
]

ADVANCED_COLUMNS = [
    "category_codes",
    "service_codes",
    "address",
    "how_to_buy",
    "size",
    "results",
    "quote",
    "logo",
    "logo_medium",
]


# =========================
# Helpers
# =========================

def _json_files(dataset_dir: Path) -> List[Path]:
    if not dataset_dir.exists():
        return []
    return sorted(p for p in dataset_dir.iterdir() if p.suffix == ".json")


def load_dataset_items(dataset_dir: Path) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    for fp in _json_files(dataset_dir):
        try:
            items.append(json.loads(fp.read_text(encoding="utf-8")))
        except Exception:
            continue
    return items


def _ensure_https(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if u.startswith("http://") or u.startswith("https://"):
        return u
    return "https://" + u.lstrip("/")


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return _clean_text(v) if isinstance(v, str) else v
    if isinstance(v, list):
        return "; ".join(_clean_text(x) for x in v if x is not None and str(x).strip() != "")
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return _clean_text(v)


def _prefer_codes(row: Dict[str, Any], codes_key: str, raw_key: str, kind: str) -> list[str]:
    """
    Prefer already-extracted codes (category_codes/service_codes) if present.
    Otherwise parse from raw fields (categories/services).

    kind: "rpc" or "rss"
    """
    # 1) use explicit codes if provided by crawler
    v_codes = row.get(codes_key)
    codes = _parse_codes_any(v_codes)

    # 2) fallback: parse raw field
    if not codes:
        codes = _parse_codes_any(row.get(raw_key))

    # 3) filter by kind
    if kind == "rpc":
        return [c for c in codes if c.startswith("rpc")]
    return [c for c in codes if c.startswith("rss")]


# =========================
# Enrichment
# =========================

def enrich_items(
    items: List[Dict[str, Any]],
    category_map: Dict[str, str],
    service_map: Dict[str, str],
) -> List[Dict[str, Any]]:
    enriched: List[Dict[str, Any]] = []

    # normalize maps once (lower keys, clean labels, drop "…" etc.)
    cat_map = _normalize_tax_map(category_map or {})
    srv_map = _normalize_tax_map(service_map or {})

    for it in items:
        row = dict(it)

        # -------------------------
        # Normalize / clean fields
        # -------------------------
        row["entity_name"] = _clean_text(row.get("entity_name"))
        row["location"] = _clean_text(row.get("location"))
        row["address"] = _clean_text(row.get("address"))
        row["email"] = _clean_text(row.get("email"))
        row["phone"] = _clean_text(row.get("phone"))
        row["how_to_buy"] = _clean_text(row.get("how_to_buy"))
        row["size"] = _clean_text(row.get("size"))
        row["results"] = _clean_text(row.get("results"))
        row["quote"] = _clean_text(row.get("quote"))

        row["logo"] = _ensure_https(_clean_text(row.get("logo")))
        row["logo_medium"] = _ensure_https(_clean_text(row.get("logo_medium")))

        row["website"] = _ensure_https(_clean_text(row.get("website")))
        row["profile_url"] = _ensure_https(_clean_text(row.get("profile_url")))

        rpc_codes_from_products = map_product_labels_to_rpc(row.get("products"))
        row["products_rpc_codes"] = "; ".join(rpc_codes_from_products)
        row["products_rpc_names"] = "; ".join(cat_map.get(c, c) for c in rpc_codes_from_products)
        # -------------------------
        # Codes (prefer crawler output)
        # -------------------------
        cat_codes = _prefer_codes(row, "category_codes", "categories", kind="rpc")
        srv_codes = _prefer_codes(row, "service_codes", "services", kind="rss")

        # normalize ssN -> rssN (some pages show ss1)
        srv_codes2: list[str] = []
        seen_srv: set[str] = set()
        for c in (srv_codes or []):
            cc = (c or "").strip().lower()
            if not cc:
                continue
            if cc.startswith("ss") and cc[2:].isdigit():
                cc = "rss" + cc[2:]
            # keep only valid rssN
            if cc.startswith("rss") and cc[3:].isdigit():
                if cc not in seen_srv:
                    seen_srv.add(cc)
                    srv_codes2.append(cc)
        srv_codes = srv_codes2

        # also normalize cats (safety)
        cat_codes2: list[str] = []
        seen_cat: set[str] = set()
        for c in (cat_codes or []):
            cc = (c or "").strip().lower()
            if cc.startswith("rpc") and cc[3:].isdigit():
                if cc not in seen_cat:
                    seen_cat.add(cc)
                    cat_codes2.append(cc)
        cat_codes = cat_codes2

        # store codes as semicolon strings
        row["category_codes"] = "; ".join(cat_codes)
        row["service_codes"] = "; ".join(srv_codes)

        # -------------------------
        # Decode -> names
        # -------------------------
        cat_names = _decode_codes(cat_codes, cat_map)
        srv_names = _decode_codes(srv_codes, srv_map)

        row["category_names"] = "; ".join(cat_names)
        row["service_names"] = "; ".join(srv_names)

        # -------------------------
        # Products normalization
        # -------------------------
        if isinstance(row.get("products"), list):
            row["products"] = "; ".join(_clean_text(x) for x in row["products"] if _clean_text(x))
        else:
            row["products"] = _clean_text(row.get("products"))

        enriched.append(row)

    return enriched

# =========================
# Writers
# =========================

def write_csv_file(items: List[Dict[str, Any]], path: Path, columns: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for it in items:
            w.writerow({k: _cell(it.get(k)) for k in columns})


def write_xlsx_file(items: List[Dict[str, Any]], path: Path, columns: List[str]) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(columns)
    for it in items:
        ws.append([_cell(it.get(k)) for k in columns])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# =========================
# Public API
# =========================

def export_outputs(
    dataset_dir: Path,
    out_base: Path,
    write_csv: bool = True,
    write_xlsx: bool = True,
    columns_mode: str = "default",
    category_map: Optional[Dict[str, str]] = None,
    service_map: Optional[Dict[str, str]] = None,
) -> Dict[str, str]:
    items = load_dataset_items(dataset_dir)
    items = enrich_items(items, category_map or {}, service_map or {})

    columns = DEFAULT_COLUMNS if columns_mode != "all" else DEFAULT_COLUMNS + ADVANCED_COLUMNS
    paths: Dict[str, str] = {}

    if write_csv:
        p = out_base.with_suffix(".csv")
        write_csv_file(items, p, columns)
        paths["csv"] = str(p)

    if write_xlsx:
        p = out_base.with_suffix(".xlsx")
        write_xlsx_file(items, p, columns)
        paths["xlsx"] = str(p)

    return paths
